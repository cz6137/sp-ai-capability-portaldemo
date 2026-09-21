"""生成明确标注为合成数据的中文浏览器回归文件，不使用业务资料。"""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
import wave

base = Path(__file__).resolve().parents[1] / 'outputs/Chrome验收-20260905/中文 测试材料'
base.mkdir(parents=True, exist_ok=False)
markdown = '# 中文格式回归（合成样例）\n\n这份文件仅用于软件测试，不是真实项目记录。\n\n## 内容检查\n\n保留**加粗**、*斜体*、中文标点。\n\n1. 第一项\n2. 第二项\n\n| 检查项 | 结果 |\n| --- | --- |\n| 中文字符 | 待核对 |\n| 金额 123.45 元 | 待核对 |\n\n```text\n第一行\n第二行\n```\n'
(base / '中文 格式回归.md').write_text(markdown, encoding='utf-8')
doc = Document()
doc.styles['Normal'].font.name = 'Microsoft YaHei'
doc.styles['Normal'].font.size = Pt(12)
doc.styles['Normal'].element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
doc.add_heading('中文 Word 回归（合成样例）', 0)
doc.add_paragraph('本文件仅用于测试，不是真实项目记录。')
doc.add_heading('一、文字与表格', 1)
p = doc.add_paragraph('中文标点：金额 123.45 元；编号 A-001。'); p.add_run('加粗重点').bold = True
table = doc.add_table(rows=3, cols=2)
for row, values in zip(table.rows, [('检查项', '预期'), ('普通中文', '完整保留'), ('跨页内容', '人工核对')]):
    for cell, value in zip(row.cells, values): cell.text = value
doc.add_paragraph('第一条事项', 'List Number'); doc.add_paragraph('第二条事项', 'List Number')
doc.add_page_break(); doc.add_heading('二、第二页', 1); doc.add_paragraph('第二页验收标记：中文内容不能丢失。')
doc.save(base / '中文 表格分页.docx')
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
c = canvas.Canvas(str(base / '中文 文本两页.pdf'))
for page in (1, 2):
    c.setFont('STSong-Light', 18); c.drawString(50, 780, f'中文文本回归 第{page}页（合成样例）')
    c.setFont('STSong-Light', 12); c.drawString(50, 740, '本文件仅用于软件测试，金额 123.45 元，标点与中文应保留。')
    c.drawString(50, 700, f'第{page}页独立验收标记：不得丢失文字内容。'); c.showPage()
c.save()
c = canvas.Canvas(str(base / '无文字图形.pdf')); c.rect(50, 50, 300, 500); c.showPage(); c.save()
(base / '损坏文件.docx').write_text('这不是一个 Word 文件', encoding='utf-8')
(base / '损坏文件.pdf').write_bytes(b'not a pdf')
(base / '空文件.md').write_bytes(b'')
with wave.open(str(base / '合成静音.wav'), 'wb') as w:
    w.setnchannels(1); w.setsampwidth(2); w.setframerate(16000); w.writeframes(b'\0' * 32000)
print(base)
